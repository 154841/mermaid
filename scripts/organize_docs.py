#!/usr/bin/env python3
"""
organize_docs.py - Automatically categorize and organize Word documents.

Scans a directory tree for .doc/.docx files, analyzes their content,
categorizes them by type (contracts, invoices, resumes, reports, etc.),
and copies/moves them into an organized folder structure.

Requirements: python-docx (pip install python-docx)
Optional: antiword (for legacy .doc files, install via: apt install antiword)

Usage:
    python organize_docs.py                        # Scan ~/Documents, copy to ~/Documents/Organized
    python organize_docs.py -s ~/Downloads -n      # Dry run on Downloads folder
    python organize_docs.py --move --verbose        # Move files with debug logging
"""

import argparse
import hashlib
import json
import logging
import os
import re
import shutil
import subprocess
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Optional

try:
    from docx import Document as DocxDocument

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# ---------------------------------------------------------------------------
# Category definitions
# ---------------------------------------------------------------------------

CATEGORIES: dict[str, dict] = {
    "contracts": {
        "keywords": [
            "agreement",
            "contract",
            "hereby",
            "parties",
            "terms and conditions",
            "whereas",
            "obligations",
            "liability",
            "indemnify",
            "governing law",
            "termination",
            "breach",
            "clause",
            "binding",
            "executed",
            "witnesseth",
            "covenant",
            "arbitration",
            "jurisdiction",
        ],
        "weight_boost": 1.0,
    },
    "invoices": {
        "keywords": [
            "invoice",
            "bill to",
            "ship to",
            "amount due",
            "payment terms",
            "subtotal",
            "tax",
            "total due",
            "purchase order",
            "qty",
            "unit price",
            "net amount",
            "due date",
            "remittance",
            "account number",
            "billing",
        ],
        "weight_boost": 1.2,
    },
    "resumes": {
        "keywords": [
            "experience",
            "education",
            "skills",
            "objective",
            "references",
            "employment",
            "curriculum vitae",
            "qualifications",
            "proficient",
            "certification",
            "gpa",
            "bachelor",
            "master",
            "linkedin",
            "work history",
            "accomplishments",
            "career",
        ],
        "weight_boost": 1.0,
    },
    "reports": {
        "keywords": [
            "report",
            "summary",
            "findings",
            "analysis",
            "conclusion",
            "recommendation",
            "executive summary",
            "methodology",
            "results",
            "appendix",
            "figure",
            "table",
            "abstract",
            "introduction",
            "discussion",
            "overview",
            "assessment",
        ],
        "weight_boost": 0.8,
    },
    "letters": {
        "keywords": [
            "dear",
            "sincerely",
            "regards",
            "to whom it may concern",
            "yours faithfully",
            "enclosed",
            "please find",
            "cordially",
            "thank you for",
            "i am writing",
            "best regards",
            "yours truly",
        ],
        "weight_boost": 1.1,
    },
    "technical": {
        "keywords": [
            "specification",
            "architecture",
            "implementation",
            "api",
            "database",
            "algorithm",
            "configuration",
            "deployment",
            "requirements",
            "system design",
            "protocol",
            "interface",
            "documentation",
            "version",
            "module",
            "framework",
            "integration",
            "endpoint",
            "schema",
        ],
        "weight_boost": 0.9,
    },
    "presentations": {
        "keywords": [
            "slide",
            "presentation",
            "agenda",
            "overview",
            "key takeaways",
            "next steps",
            "questions",
            "outline",
            "objectives",
            "highlights",
            "roadmap",
        ],
        "weight_boost": 1.0,
    },
    "legal": {
        "keywords": [
            "plaintiff",
            "defendant",
            "court",
            "statute",
            "legal",
            "attorney",
            "counsel",
            "motion",
            "filing",
            "judgment",
            "testimony",
            "deposition",
            "subpoena",
            "affidavit",
            "complaint",
            "petition",
        ],
        "weight_boost": 1.1,
    },
    "financial": {
        "keywords": [
            "revenue",
            "profit",
            "loss",
            "balance sheet",
            "cash flow",
            "budget",
            "forecast",
            "quarterly",
            "fiscal",
            "earnings",
            "dividend",
            "assets",
            "liabilities",
            "equity",
            "depreciation",
            "amortization",
            "audit",
        ],
        "weight_boost": 1.0,
    },
}

DEFAULT_CATEGORY = "uncategorized"

SKIP_DIRS = {
    ".git",
    "node_modules",
    "__pycache__",
    ".venv",
    "venv",
    ".tox",
    ".mypy_cache",
    ".pytest_cache",
}

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------


def setup_logging(log_file: Optional[Path], verbose: bool) -> logging.Logger:
    """Configure logging to console and optional file."""
    logger = logging.getLogger("organize_docs")
    logger.setLevel(logging.DEBUG if verbose else logging.INFO)

    console = logging.StreamHandler(sys.stdout)
    console.setLevel(logging.DEBUG if verbose else logging.INFO)
    console.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
    logger.addHandler(console)

    if log_file:
        log_file.parent.mkdir(parents=True, exist_ok=True)
        fh = logging.FileHandler(log_file)
        fh.setLevel(logging.DEBUG)
        fh.setFormatter(
            logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
        )
        logger.addHandler(fh)

    return logger


# ---------------------------------------------------------------------------
# File discovery
# ---------------------------------------------------------------------------


def find_documents(
    search_dir: Path,
    output_dir: Path,
    extensions: tuple[str, ...] = (".docx",),
) -> list[Path]:
    """Recursively find Word documents, skipping hidden/irrelevant dirs."""
    results: list[Path] = []
    output_resolved = output_dir.resolve()

    for root, dirs, files in os.walk(search_dir):
        # Skip hidden directories and known non-document directories
        dirs[:] = [
            d for d in dirs if not d.startswith(".") and d not in SKIP_DIRS
        ]

        root_path = Path(root)
        if root_path.resolve() == output_resolved or (
            output_resolved != search_dir.resolve()
            and str(root_path.resolve()).startswith(str(output_resolved))
        ):
            dirs.clear()
            continue

        for fname in files:
            # Skip Word temp files (~$filename.docx)
            if fname.startswith("~$"):
                continue
            if any(fname.lower().endswith(ext) for ext in extensions):
                results.append(root_path / fname)

    results.sort(key=lambda p: p.name.lower())
    return results


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------


def extract_text_docx(file_path: Path) -> tuple[str, dict]:
    """Extract full text and metadata from a .docx file."""
    doc = DocxDocument(str(file_path))

    paragraphs: list[str] = []

    # Body paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text)

    # Table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)

    # Headers and footers
    for section in doc.sections:
        for part in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            try:
                if part and not part.is_linked_to_previous:
                    for p in part.paragraphs:
                        if p.text.strip():
                            paragraphs.append(p.text)
            except Exception:
                pass

    full_text = "\n".join(paragraphs)

    # Metadata
    core = doc.core_properties
    metadata = {
        "author": core.author or "",
        "title": core.title or "",
        "subject": core.subject or "",
        "keywords": core.keywords or "",
        "category": core.category or "",
        "created": str(core.created) if core.created else "",
        "modified": str(core.modified) if core.modified else "",
    }

    return full_text, metadata


def extract_text_doc(file_path: Path) -> tuple[str, dict]:
    """Best-effort extraction from legacy .doc binary format."""
    metadata = {
        "author": "",
        "title": "",
        "subject": "",
        "keywords": "",
        "category": "",
        "created": "",
        "modified": "",
    }

    # Try antiword first
    try:
        result = subprocess.run(
            ["antiword", str(file_path)],
            capture_output=True,
            text=True,
            timeout=30,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout, metadata
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # Fallback: extract ASCII strings from binary
    try:
        with open(file_path, "rb") as f:
            raw = f.read()
        text_chunks = re.findall(rb"[\x20-\x7e]{4,}", raw)
        text = " ".join(chunk.decode("ascii") for chunk in text_chunks)
        return text, metadata
    except Exception:
        return "", metadata


# ---------------------------------------------------------------------------
# Categorization
# ---------------------------------------------------------------------------


def categorize_document(
    text: str, metadata: dict, min_score: int = 2
) -> tuple[str, dict[str, float]]:
    """Score document against all categories and return the best match."""
    text_lower = text.lower()
    scores: dict[str, float] = {}

    for cat_name, cat_config in CATEGORIES.items():
        score = 0.0
        keywords = cat_config["keywords"]
        boost = cat_config.get("weight_boost", 1.0)

        for kw in keywords:
            pattern = r"\b" + re.escape(kw) + r"\b"
            matches = len(re.findall(pattern, text_lower))
            score += matches

        # Metadata bonus (2x weight)
        meta_text = " ".join(
            [
                metadata.get("title", ""),
                metadata.get("subject", ""),
                metadata.get("keywords", ""),
                metadata.get("category", ""),
            ]
        ).lower()

        if meta_text.strip():
            for kw in keywords:
                if kw in meta_text:
                    score += 2

        scores[cat_name] = score * boost

    if not scores:
        return DEFAULT_CATEGORY, scores

    best = max(scores, key=lambda k: scores[k])
    if scores[best] < min_score:
        return DEFAULT_CATEGORY, scores

    return best, scores


# ---------------------------------------------------------------------------
# File operations
# ---------------------------------------------------------------------------


def compute_file_hash(file_path: Path) -> str:
    """Compute SHA-256 hash for duplicate detection."""
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def generate_safe_filename(src: Path, dest_dir: Path) -> Optional[Path]:
    """Return a safe destination path, or None if file is an exact duplicate."""
    dest = dest_dir / src.name
    if not dest.exists():
        return dest

    # Exact duplicate check
    if compute_file_hash(src) == compute_file_hash(dest):
        return None

    # Name collision — append counter
    stem = src.stem
    suffix = src.suffix
    counter = 1
    while True:
        dest = dest_dir / f"{stem}_{counter}{suffix}"
        if not dest.exists():
            return dest
        if compute_file_hash(src) == compute_file_hash(dest):
            return None
        counter += 1


def organize_file(
    src: Path,
    output_dir: Path,
    category: str,
    move: bool,
    dry_run: bool,
    logger: logging.Logger,
) -> dict:
    """Copy or move a single file into the category folder."""
    cat_dir = output_dir / category
    dest = generate_safe_filename(src, cat_dir)

    operation = "move" if move else "copy"
    result: dict = {
        "source": str(src),
        "category": category,
        "operation": operation,
    }

    if dest is None:
        result["status"] = "skipped"
        result["reason"] = "duplicate"
        logger.info(f"Skipping duplicate: {src.name}")
        return result

    result["destination"] = str(dest)

    if dry_run:
        result["status"] = "dry_run"
        logger.info(f"[DRY RUN] Would {operation} {src} -> {dest}")
        return result

    cat_dir.mkdir(parents=True, exist_ok=True)
    if move:
        shutil.move(str(src), str(dest))
    else:
        shutil.copy2(str(src), str(dest))

    result["status"] = "success"
    action = "Moved" if move else "Copied"
    logger.info(f"{action}: {src.name} -> {category}/")
    return result


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------


def generate_report(
    operations: list[dict],
    output_dir: Path,
    dry_run: bool,
    logger: logging.Logger,
) -> str:
    """Generate a summary report as text and JSON."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    by_category: dict[str, list[dict]] = defaultdict(list)
    total = len(operations)
    success = 0
    skipped = 0
    errors = 0

    for op in operations:
        status = op.get("status", "")
        cat = op.get("category", DEFAULT_CATEGORY)
        by_category[cat].append(op)
        if status in ("success", "dry_run"):
            success += 1
        elif status == "skipped":
            skipped += 1
        elif status == "error":
            errors += 1

    lines = [
        "=" * 60,
        "Document Organization Report",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Mode: {'DRY RUN' if dry_run else 'LIVE'}",
        "=" * 60,
        "",
        f"Total files found:      {total}",
        f"Successfully organized:  {success}",
        f"Skipped (duplicates):   {skipped}",
        f"Errors:                 {errors}",
        "",
        "-" * 40,
        "Breakdown by Category:",
        "-" * 40,
    ]

    for cat in sorted(by_category.keys()):
        items = by_category[cat]
        lines.append(f"\n  {cat.upper()} ({len(items)} files):")
        for item in items:
            src = Path(item["source"]).name
            status = item["status"]
            lines.append(f"    - {src} [{status}]")

    error_items = [op for op in operations if op.get("status") == "error"]
    if error_items:
        lines.append(f"\n{'=' * 40}")
        lines.append("ERRORS:")
        for item in error_items:
            lines.append(
                f"  {item['source']}: {item.get('reason', 'unknown')}"
            )

    lines.append(f"\n{'=' * 60}")
    report_text = "\n".join(lines)

    if not dry_run:
        output_dir.mkdir(parents=True, exist_ok=True)
        report_txt = output_dir / f"organization_report_{timestamp}.txt"
        report_json = output_dir / f"organization_report_{timestamp}.json"

        report_txt.write_text(report_text)

        json_data = {
            "timestamp": timestamp,
            "dry_run": dry_run,
            "total": total,
            "success": success,
            "skipped": skipped,
            "errors": errors,
            "operations": operations,
        }
        report_json.write_text(json.dumps(json_data, indent=2))

        logger.info(f"Report written to: {report_txt}")

    return report_text


# ---------------------------------------------------------------------------
# Processing pipeline
# ---------------------------------------------------------------------------


def process_file(
    file_path: Path,
    output_dir: Path,
    move: bool,
    dry_run: bool,
    include_doc: bool,
    min_score: int,
    logger: logging.Logger,
) -> Optional[dict]:
    """Process a single document: extract, categorize, organize."""
    try:
        ext = file_path.suffix.lower()

        if ext == ".docx":
            text, metadata = extract_text_docx(file_path)
        elif ext == ".doc":
            if not include_doc:
                logger.debug(f"Skipping .doc file: {file_path}")
                return {
                    "source": str(file_path),
                    "status": "skipped",
                    "reason": "legacy .doc (use --include-doc)",
                    "category": "",
                }
            text, metadata = extract_text_doc(file_path)
        else:
            return None

        if not text.strip():
            logger.warning(f"No text extracted: {file_path.name}")
            category = DEFAULT_CATEGORY
            scores: dict[str, float] = {}
        else:
            category, scores = categorize_document(text, metadata, min_score)
            logger.debug(f"Scores for {file_path.name}: {scores}")

        logger.info(f"  -> Category: {category}")
        return organize_file(
            file_path, output_dir, category, move, dry_run, logger
        )

    except PermissionError:
        logger.error(f"Permission denied: {file_path}")
        return {
            "source": str(file_path),
            "status": "error",
            "reason": "permission denied",
            "category": "",
        }
    except FileNotFoundError:
        logger.error(f"File not found: {file_path}")
        return {
            "source": str(file_path),
            "status": "error",
            "reason": "file not found",
            "category": "",
        }
    except Exception as e:
        logger.error(
            f"Error processing {file_path}: {type(e).__name__}: {e}"
        )
        return {
            "source": str(file_path),
            "status": "error",
            "reason": f"{type(e).__name__}: {e}",
            "category": "",
        }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Organize Word documents by content category.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  %(prog)s                              # Scan ~/Documents, copy to ~/Documents/Organized
  %(prog)s -s ~/Downloads -n            # Dry run on Downloads folder
  %(prog)s --move --verbose             # Move files with debug logging
  %(prog)s -s /path/to/docs -o /path/to/organized --min-score 3
        """,
    )
    home = Path.home()
    parser.add_argument(
        "--source",
        "-s",
        type=Path,
        default=home / "Documents",
        help="Directory to scan (default: ~/Documents)",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=Path,
        default=home / "Documents" / "Organized",
        help="Output directory (default: ~/Documents/Organized)",
    )
    parser.add_argument(
        "--move",
        action="store_true",
        help="Move files instead of copying (default: copy)",
    )
    parser.add_argument(
        "--dry-run",
        "-n",
        action="store_true",
        help="Preview changes without executing",
    )
    parser.add_argument(
        "--verbose",
        "-v",
        action="store_true",
        help="Enable debug logging",
    )
    parser.add_argument(
        "--log-file",
        type=Path,
        default=None,
        help="Path to log file",
    )
    parser.add_argument(
        "--min-score",
        type=int,
        default=2,
        help="Minimum category score threshold (default: 2)",
    )
    parser.add_argument(
        "--include-doc",
        action="store_true",
        help="Also process legacy .doc files (requires antiword)",
    )
    return parser.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main() -> None:
    args = parse_args()

    if not HAS_DOCX:
        print("ERROR: python-docx is required. Install it with:")
        print("  pip install python-docx")
        sys.exit(1)

    source = args.source.expanduser().resolve()
    output = args.output.expanduser().resolve()

    if not source.exists():
        print(f"ERROR: Source directory does not exist: {source}")
        sys.exit(1)

    if not source.is_dir():
        print(f"ERROR: Source is not a directory: {source}")
        sys.exit(1)

    # Logging
    log_file = args.log_file
    if log_file is None and not args.dry_run:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        log_file = output / f"organize_log_{timestamp}.log"
    logger = setup_logging(log_file, args.verbose)

    # Extensions
    extensions = [".docx"]
    if args.include_doc:
        extensions.append(".doc")

    logger.info(f"Source directory: {source}")
    logger.info(f"Output directory: {output}")
    logger.info(f"Mode: {'MOVE' if args.move else 'COPY'}")
    logger.info(f"Dry run: {args.dry_run}")
    logger.info(f"Extensions: {extensions}")

    # Discover files
    files = find_documents(source, output, tuple(extensions))
    logger.info(f"Found {len(files)} document(s)")

    if not files:
        logger.info("No documents found. Nothing to do.")
        return

    # Process
    operations: list[dict] = []
    for i, file_path in enumerate(files, 1):
        logger.info(f"[{i}/{len(files)}] Processing: {file_path.name}")
        result = process_file(
            file_path,
            output,
            args.move,
            args.dry_run,
            args.include_doc,
            args.min_score,
            logger,
        )
        if result:
            operations.append(result)

    # Report
    report = generate_report(operations, output, args.dry_run, logger)
    print(report)


if __name__ == "__main__":
    main()
